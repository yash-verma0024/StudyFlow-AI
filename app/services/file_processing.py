from pathlib import Path
import json
import os
from typing import Tuple

from app.config import ALLOWED_EXTENSIONS, MAX_FILE_SIZE


def validate_upload(file_name: str, content_length: int | None) -> None:
    if not file_name:
        raise ValueError("No file selected.")

    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Please upload PDF, DOCX, PPTX, or image files.")

    if content_length is not None and content_length > MAX_FILE_SIZE:
        raise ValueError("File is too large. Please upload a file smaller than 15MB.")


def extract_text(file_path: str) -> Tuple[str, str]:
    ext = Path(file_path).suffix.lower().lstrip(".")

    if ext == "pdf":
        return extract_pdf_text(file_path)
    if ext == "docx":
        return extract_docx_text(file_path)
    if ext == "pptx":
        return extract_pptx_text(file_path)
    if ext in {"png", "jpg", "jpeg"}:
        return extract_image_text(file_path)

    raise ValueError("Unsupported file type.")


def extract_pdf_text(file_path: str) -> Tuple[str, str]:
    try:
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            pages.append(page.extract_text() or "")
        text = "\n\n".join(pages).strip()
        if not text:
            raise ValueError("Unable to extract readable text from the PDF.")
        return text, "pdf"
    except Exception as exc:
        raise ValueError(f"Failed to extract PDF text: {exc}") from exc


def extract_docx_text(file_path: str) -> Tuple[str, str]:
    try:
        from docx import Document

        document = Document(file_path)
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        text = "\n".join(paragraphs).strip()
        if not text:
            raise ValueError("Unable to extract readable text from the DOCX file.")
        return text, "docx"
    except Exception as exc:
        raise ValueError(f"Failed to extract DOCX text: {exc}") from exc


def extract_pptx_text(file_path: str) -> Tuple[str, str]:
    try:
        from pptx import Presentation

        presentation = Presentation(file_path)
        slides = []
        for slide in presentation.slides:
            text_chunks = []
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text = shape.text.strip()
                    if text:
                        text_chunks.append(text)
            slides.append("\n".join(text_chunks))

        text = "\n\n".join(slides).strip()
        if not text:
            raise ValueError("Unable to extract readable text from the PPTX file.")
        return text, "pptx"
    except Exception as exc:
        raise ValueError(f"Failed to extract PPTX text: {exc}") from exc


def extract_image_text(file_path: str) -> Tuple[str, str]:
    try:
        # Use OCR only if Google GenAI provides a reliable alternative later. For now, return a friendly failure.
        raise ValueError("Image OCR is not available in the current build. Please upload a PDF, DOCX, or PPTX file for the guaranteed demo path.")
    except Exception:
        raise
